"""對 Firebase 模擬器執行完整 API 流程測試（一次跑完，不重複）。
用法：python tests/api_flow_test.py http://127.0.0.1:5055
"""
import base64
import io
import os
import sys
from datetime import datetime, timedelta, timezone

import requests

BASE = (sys.argv[1] if len(sys.argv) > 1 else "http://127.0.0.1:5055").rstrip("/")
API = BASE + "/api"
AUTH = "http://127.0.0.1:9099/identitytoolkit.googleapis.com/v1/accounts:signUp?key=fake-api-key"
PNG = "data:image/png;base64," + base64.b64encode(
    base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")).decode()
results = []


def check(name, cond, detail=""):
    results.append((name, bool(cond)))
    print(("PASS " if cond else "FAIL ") + name + (f"  [{detail}]" if detail and not cond else ""))


def docx_b64():
    from docx import Document
    d = Document()
    d.add_paragraph("第一條　合作範圍：乙方提供軟體服務。")
    d.add_paragraph("第二條　保密義務：雙方應保密三年。")
    buf = io.BytesIO()
    d.save(buf)
    return base64.b64encode(buf.getvalue()).decode()


tok = requests.post(AUTH, json={"returnSecureToken": True}).json()["idToken"]
H = {"Authorization": "Bearer " + tok}
exp = (datetime.now(timezone.utc) + timedelta(days=7)).isoformat()

# ---- 三方合約：上傳、翻譯、封面、簽署、拒絕後鎖定
r = requests.post(API + "/contracts", headers=H, json={
    "title": "全球服務合作合約", "summary": "雙方就軟體服務合作達成協議。", "body": "第一條　合作範圍。\n第二條　服務費用。",
    "language": "zh-TW", "coverId": "business-gold",
    "parties": [{"role": "甲方", "email": "a@example.com"}, {"role": "乙方", "email": "b@example.com"}, {"role": "丙方", "email": "c@example.com"}]})
check("建立三方合約", r.status_code == 200, r.text)
t, c = r.json()["tenantId"], r.json()["contractId"]
check("未登入不可建立合約", requests.post(API + "/contracts", json={"title": "x"}).status_code == 401)
r = requests.post(API + "/contracts/document", headers=H, json={"tenantId": t, "contractId": c, "fileName": "服務合約.docx",
    "contentType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "dataBase64": docx_b64()})
check("上傳 Word 文件並擷取文字", r.status_code == 200 and r.json().get("textLength", 0) > 10, r.text)
r = requests.post(API + "/contracts/document", headers=H, json={"tenantId": t, "contractId": c, "fileName": "a.exe", "contentType": "application/x-msdownload", "dataBase64": "AAAA"})
check("拒絕不支援的檔案類型", r.status_code == 400 and r.json()["error"] == "unsupported_file_type", r.text)
r = requests.get(API + f"/contracts/document?tenantId={t}&contractId={c}", headers=H)
check("建立者可下載文件", r.status_code == 200 and r.content[:2] == b"PK", r.status_code)
r = requests.post(API + "/contracts/translate", headers=H, json={"tenantId": t, "contractId": c, "lang": "zh-CN"})
check("翻譯為簡體（含文件文字）", r.status_code == 200 and "协议" in r.json()["fields"]["summary"] and "保密义务" in r.json()["fields"]["documentText"], r.text[:200])
r = requests.post(API + "/contracts/translate", headers=H, json={"tenantId": t, "contractId": c, "lang": "en"})
check("翻譯為英文（模擬器以假翻譯代替 Cloud Translation）", r.status_code == 200 and r.json()["fields"]["title"].startswith("[en]"), r.text[:200])
st = requests.get(API + f"/contracts/status?tenantId={t}&contractId={c}", headers=H).json()
check("封面選擇已保存", st["contract"]["coverId"] == "business-gold", st["contract"].get("coverId"))

r = requests.post(API + "/signature/invite", headers=H, json={"tenantId": t, "contractId": c, "expiresAt": exp})
check("發送三方簽署邀請", r.status_code == 200 and len(r.json()["invites"]) == 3, r.text[:200])
links = [v["previewUrl"] for v in r.json()["invites"]]
check("簽署連結格式 /sign/<token>", all("/sign/" in u for u in links), links[0])
toks = [u.split("/sign/")[1] for u in links]
check("不可重複發送邀請", requests.post(API + "/signature/invite", headers=H, json={"tenantId": t, "contractId": c, "expiresAt": exp}).status_code == 409)
r = requests.post(API + "/contracts/document", headers=H, json={"tenantId": t, "contractId": c, "fileName": "x.txt", "contentType": "text/plain", "dataBase64": base64.b64encode(b"x").decode()})
check("發出邀請後文件鎖定", r.status_code == 409, r.text)

r = requests.get(API + "/signature/invite", params={"token": toks[0], "lang": "en"})
check("簽署人讀取合約（英文）", r.status_code == 200 and r.json()["contract"]["translation"]["lang"] == "en" and r.json()["contract"]["coverId"] == "business-gold", r.text[:200])
check("簽署人可開啟文件", requests.get(API + "/signature/document", params={"token": toks[0]}).status_code == 200)
r = requests.post(API + "/signature/action", json={"token": toks[0], "action": "sign", "name": "王小明", "consent": True})
check("沒有手寫簽名不能簽", r.status_code == 400 and r.json()["error"] == "signature_required", r.text)
r = requests.post(API + "/signature/action", json={"token": toks[0], "action": "sign", "name": "王小明", "consent": True, "signatureImage": PNG})
check("甲方簽署", r.status_code == 200 and r.json()["status"] == "partially_signed", r.text)
r = requests.post(API + "/signature/action", json={"token": toks[0], "action": "sign", "name": "王小明", "consent": True, "signatureImage": PNG})
check("連結用過即失效", r.status_code == 409 and r.json()["error"] == "invite_already_consumed", r.text)
check("用過的連結無法再讀取", requests.get(API + "/signature/invite", params={"token": toks[0]}).status_code == 409)
r = requests.post(API + "/signature/action", json={"token": toks[1], "action": "decline", "name": "李大華"})
check("拒絕必須填原因", r.status_code == 400 and r.json()["error"] == "reason_required", r.text)
r = requests.post(API + "/signature/action", json={"token": toks[1], "action": "decline", "name": "李大華", "reason": "價格不同意"})
check("乙方拒絕", r.status_code == 200 and r.json()["status"] == "declined", r.text)
r = requests.post(API + "/signature/action", json={"token": toks[2], "action": "sign", "name": "陳丙", "consent": True, "signatureImage": PNG})
check("拒絕後丙方不能再簽（修正後端漏洞）", r.status_code == 409 and r.json()["error"] == "contract_closed", r.text)
st = requests.get(API + f"/contracts/status?tenantId={t}&contractId={c}", headers=H).json()
check("合約維持已被拒絕", st["contract"]["status"] == "declined", st["contract"]["status"])
check("各方狀態正確", [v["status"] for v in st["invites"]] == ["signed", "declined", "pending"], [v["status"] for v in st["invites"]])

tok2 = requests.post(AUTH, json={"returnSecureToken": True}).json()["idToken"]
check("其他使用者看不到這份合約", requests.get(API + f"/contracts/status?tenantId={t}&contractId={c}", headers={"Authorization": "Bearer " + tok2}).status_code == 403)

# ---- 兩方合約：全數簽署
r = requests.post(API + "/contracts", headers=H, json={"title": "Service Agreement", "summary": "Both parties agree to the terms.", "language": "en", "coverId": "tech-blue",
    "parties": [{"role": "Party A", "email": "x@example.com"}, {"role": "Party B", "email": "y@example.com"}]})
t2, c2 = r.json()["tenantId"], r.json()["contractId"]
inv = requests.post(API + "/signature/invite", headers=H, json={"tenantId": t2, "contractId": c2, "expiresAt": exp}).json()["invites"]
r = requests.get(API + "/signature/invite", params={"token": inv[0]["previewUrl"].split("/sign/")[1], "lang": "zh-CN"})
check("英文合約翻成簡體", r.status_code == 200 and "translation" in r.json()["contract"], r.text[:200])
out = [requests.post(API + "/signature/action", json={"token": v["previewUrl"].split("/sign/")[1], "action": "sign", "name": f"Signer {i}", "consent": True, "signatureImage": PNG}).json() for i, v in enumerate(inv)]
check("兩方全數簽署", out[-1].get("status") == "fully_signed", out)
st = requests.get(API + f"/contracts/status?tenantId={t2}&contractId={c2}", headers=H).json()
check("簽名圖路徑已記錄", all(v["status"] == "signed" for v in st["invites"]) and st["contract"]["signedCount"] == 2)

# ---- 下載已簽署 PDF
r = requests.get(API + f"/contracts/signed-pdf?tenantId={t}&contractId={c}", headers=H)
check("未全數簽署不能下載 PDF", r.status_code == 409 and r.json()["error"] == "contract_not_fully_signed", r.text[:120])
r = requests.get(API + f"/contracts/signed-pdf?tenantId={t2}&contractId={c2}", headers=H)
check("建立者下載已簽署 PDF", r.status_code == 200 and r.content[:4] == b"%PDF" and "attachment" in r.headers.get("Content-Disposition", ""), r.status_code)
open("/tmp/signed_owner.pdf", "wb").write(r.content)
tk0 = inv[0]["previewUrl"].split("/sign/")[1]
r = requests.get(API + "/signature/invite", params={"token": tk0})
check("簽完後再開連結會提示可下載", r.status_code == 409 and r.json().get("signedPdfAvailable") is True, r.text[:160])
r = requests.get(API + "/signature/signed-pdf", params={"token": tk0})
check("簽署人用原連結下載 PDF", r.status_code == 200 and r.content[:4] == b"%PDF", r.status_code)
r = requests.get(API + "/signature/signed-pdf", params={"token": toks[2]})
check("未簽署的人不能下載", r.status_code == 403, r.status_code)
check("其他使用者不能下載", requests.get(API + f"/contracts/signed-pdf?tenantId={t2}&contractId={c2}", headers={"Authorization": "Bearer " + tok2}).status_code == 403)

# ---- 中文合約 + 上傳 PDF：產出含原始 PDF 頁面
import pypdf
from fpdf import FPDF
src = FPDF(); src.add_page(); src.set_font("helvetica", size=14); src.cell(0, 10, "ORIGINAL UPLOADED PDF PAGE"); orig = bytes(src.output())
r = requests.post(API + "/contracts", headers=H, json={"title": "三方採購合約", "summary": "測試中文 PDF。", "body": "第一條　標的。\n第二條　價金。", "coverId": "crimson-seal",
    "parties": [{"role": "甲方", "email": "p1@example.com"}, {"role": "乙方", "email": "p2@example.com"}, {"role": "丙方", "email": "p3@example.com"}]})
t3, c3 = r.json()["tenantId"], r.json()["contractId"]
requests.post(API + "/contracts/document", headers=H, json={"tenantId": t3, "contractId": c3, "fileName": "原始合約.pdf", "contentType": "application/pdf", "dataBase64": base64.b64encode(orig).decode()})
inv3 = requests.post(API + "/signature/invite", headers=H, json={"tenantId": t3, "contractId": c3, "expiresAt": exp}).json()["invites"]
for i, v in enumerate(inv3):
    requests.post(API + "/signature/action", json={"token": v["previewUrl"].split("/sign/")[1], "action": "sign", "name": ["王小明", "李大華", "陳丙"][i], "consent": True, "signatureImage": PNG})
r = requests.get(API + f"/contracts/signed-pdf?tenantId={t3}&contractId={c3}", headers=H)
ok3 = r.status_code == 200 and r.content[:4] == b"%PDF"
if ok3:
    open("/tmp/signed_zh.pdf", "wb").write(r.content)
    rd = pypdf.PdfReader(io.BytesIO(r.content)); text = "".join(pg.extract_text() or "" for pg in rd.pages)
    ok3 = "王小明" in text and "陳丙" in text and "ORIGINAL UPLOADED PDF PAGE" in text and "三方採購合約" in text
check("中文 PDF 含三方簽名與原始 PDF 頁", ok3, r.status_code)
r2 = requests.get(API + f"/contracts/signed-pdf?tenantId={t3}&contractId={c3}", headers=H)
check("再次下載使用快取（內容一致）", r2.content == r.content)


# ---- 合約編號 C-YYMM001、可修改、可查詢；瀏覽所有合約；合約內容頁
import re as _re
TWN = datetime.now(timezone(timedelta(hours=8))).strftime("%y%m")
tokN = requests.post(AUTH, json={"returnSecureToken": True}).json()["idToken"]; HN = {"Authorization": "Bearer " + tokN}
P2 = [{"role": "甲方", "email": "n1@example.com"}, {"role": "乙方", "email": "n2@example.com"}]
nx = requests.get(API + "/contracts/next-number", headers=HN).json().get("contractNo")
check("新帳號預設下一號 C-YYMM001", nx == f"C-{TWN}001", nx)
a = requests.post(API + "/contracts", headers=HN, json={"title": "編號測試 A 採購", "parties": P2}).json()
b = requests.post(API + "/contracts", headers=HN, json={"title": "編號測試 B 租賃", "parties": P2}).json()
check("自動編號依序 001、002", a.get("contractNo") == f"C-{TWN}001" and b.get("contractNo") == f"C-{TWN}002", (a.get("contractNo"), b.get("contractNo")))
cu = requests.post(API + "/contracts", headers=HN, json={"title": "自訂編號合約", "parties": P2, "contractNo": "SAM-2026-01"}).json()
check("建立時可自訂編號", cu.get("contractNo") == "SAM-2026-01", cu)
r = requests.post(API + "/contracts", headers=HN, json={"title": "重複", "parties": P2, "contractNo": "sam-2026-01"})
check("編號不可重複（不分大小寫）", r.status_code == 409 and r.json()["error"] == "contract_no_taken", r.text[:100])
r = requests.post(API + "/contracts", headers=HN, json={"title": "格式錯", "parties": P2, "contractNo": "含 空白/斜線"})
check("編號格式錯誤會被拒絕", r.status_code == 400 and r.json()["error"] == "invalid_contract_no", r.text[:100])
r = requests.post(API + "/contracts", headers=HN, json={"title": "手動輸入下一號", "parties": P2, "contractNo": f"C-{TWN}005"}).json()
nx2 = requests.get(API + "/contracts/next-number", headers=HN).json().get("contractNo")
check("手動用了 005，下一個預設跳到 006", nx2 == f"C-{TWN}006", nx2)
r = requests.post(API + "/contracts/number", headers=HN, json={"contractId": a["contractId"], "contractNo": "PO-0001"})
check("事後修改編號", r.status_code == 200 and r.json()["contractNo"] == "PO-0001", r.text[:100])
r = requests.post(API + "/contracts/number", headers=HN, json={"contractId": b["contractId"], "contractNo": "PO-0001"})
check("修改成已使用的編號會被拒絕", r.status_code == 409, r.status_code)
c_new = requests.post(API + "/contracts", headers=HN, json={"title": "回收舊號", "parties": P2, "contractNo": f"C-{TWN}001"})
check("改名後舊編號可再使用", c_new.status_code == 200, c_new.text[:100])
r = requests.post(API + "/contracts/number", headers={"Authorization": "Bearer " + tok}, json={"contractId": b["contractId"], "contractNo": "HACK-1"})
check("別人不能改我的編號", r.status_code == 403, r.status_code)
st = requests.get(API + "/contracts/status", headers=HN, params={"contractNo": "po-0001"})
check("以合約編號查詢合約內容（不分大小寫）", st.status_code == 200 and st.json()["contract"]["id"] == a["contractId"] and st.json()["contract"]["contractNo"] == "PO-0001", st.text[:120])
check("查不到的編號回 404", requests.get(API + "/contracts/status", headers=HN, params={"contractNo": "NOPE-9"}).status_code == 404)
check("別人用編號查不到我的合約", requests.get(API + "/contracts/status", headers={"Authorization": "Bearer " + tok}, params={"contractNo": "PO-0001"}).status_code in (403, 404))
L = requests.get(API + "/contracts/list", headers=HN).json()
check("瀏覽所有合約：列出自己的 5 份", L.get("total") == 5 and all(x["contractNo"] for x in L["contracts"]), L.get("total"))
check("查詢：編號", [x["contractNo"] for x in requests.get(API + "/contracts/list", headers=HN, params={"q": "po-00"}).json()["contracts"]] == ["PO-0001"])
check("查詢：名稱關鍵字", [x["title"] for x in requests.get(API + "/contracts/list", headers=HN, params={"q": "租賃"}).json()["contracts"]] == ["編號測試 B 租賃"])
check("查詢：簽署方 Email", requests.get(API + "/contracts/list", headers=HN, params={"q": "n2@example"}).json()["total"] == 5)
check("篩選：狀態", requests.get(API + "/contracts/list", headers=HN, params={"status": "fully_signed"}).json()["total"] == 0)
check("別人的清單看不到我的合約", all(x["id"] not in (a["contractId"], b["contractId"]) for x in requests.get(API + "/contracts/list", headers={"Authorization": "Bearer " + tok}).json()["contracts"]))
check("未登入不能瀏覽", requests.get(API + "/contracts/list").status_code == 401)
requests.post(API + "/signature/invite", headers=HN, json={"tenantId": L["tenantId"], "contractId": a["contractId"], "expiresAt": exp})
st = requests.get(API + "/contracts/status", headers=HN, params={"contractNo": "PO-0001"}).json()
check("合約內容頁可找回待簽連結（換頁、清快取也不會遺失）", all(v.get("link", "").startswith(BASE + "/sign/") or "/sign/" in (v.get("link") or "") for v in st["invites"]), [v.get("link") for v in st["invites"]][:1])
from urllib.parse import unquote as _uq
_cd = _uq(requests.get(API + f"/contracts/signed-pdf?tenantId={t3}&contractId={c3}", headers=H).headers.get("Content-Disposition", ""))
_no3 = requests.get(API + "/contracts/status", headers=H, params={"tenantId": t3, "contractId": c3}).json()["contract"]["contractNo"]
check("PDF 檔名使用合約編號", _no3.startswith("C-") and _no3 in _cd, (_no3, _cd[:80]))
for path in ["/contracts", "/contracts/PO-0001"]:
    r = requests.get(BASE + path)
    check(f"網頁 {path} 可開啟", r.status_code == 200 and 'id="list-page"' in r.text, r.status_code)

# ---- 簽署方公司資料：公司名稱／統一編號／公司地址／代表人
PC = [{"role": "甲方", "email": "co1@example.com", "company": "台灣大哥大股份有限公司", "taxId": "04595257", "address": "台北市大安區敦化南路二段 207 號", "representative": "林之晨"},
      {"role": "乙方", "email": "co2@example.com", "company": "Acme Corp.", "taxId": "US-12-3456789", "address": "1 Market St, San Francisco", "representative": "Jane Doe"}]
rc = requests.post(API + "/contracts", headers=HN, json={"title": "公司資料合約", "parties": PC})
check("建立含公司資料的合約", rc.status_code == 200, rc.text[:120])
stc = requests.get(API + "/contracts/status", headers=HN, params={"contractId": rc.json().get("contractId")}).json()["contract"]["parties"]
check("四個欄位完整保存", stc[0]["company"] == "台灣大哥大股份有限公司" and stc[0]["taxId"] == "04595257" and stc[0]["address"].startswith("台北市") and stc[0]["representative"] == "林之晨" and stc[1]["taxId"] == "US-12-3456789", stc[0])
bad = requests.post(API + "/contracts", headers=HN, json={"title": "錯統編", "parties": [dict(PC[0], taxId="12345678"), PC[1]]})
check("台灣統編檢查碼錯誤會被拒絕", bad.status_code == 400 and bad.json()["error"] == "invalid_tax_id", bad.text[:100])
bad2 = requests.post(API + "/contracts", headers=HN, json={"title": "錯統編", "parties": [dict(PC[0], taxId="含/符號"), PC[1]]})
check("統編含非法字元會被拒絕", bad2.status_code == 400, bad2.status_code)
ok_blank = requests.post(API + "/contracts", headers=HN, json={"title": "個人簽約", "parties": P2})
check("公司欄位可留白（個人簽約）", ok_blank.status_code == 200)
check("可用公司名稱查詢", requests.get(API + "/contracts/list", headers=HN, params={"q": "台灣大哥大"}).json()["total"] == 1)
check("可用統編查詢", requests.get(API + "/contracts/list", headers=HN, params={"q": "04595257"}).json()["total"] == 1)
check("可用代表人查詢", requests.get(API + "/contracts/list", headers=HN, params={"q": "jane doe"}).json()["total"] == 1)
cidc = rc.json()["contractId"]; tnc = rc.json()["tenantId"]
inv = requests.post(API + "/signature/invite", headers=HN, json={"tenantId": tnc, "contractId": cidc, "expiresAt": exp}).json()
for v, nm in zip(inv["invites"], ["王小明", "Jane Doe"]):
    tk = v["previewUrl"].rsplit("/", 1)[1]
    pv = requests.get(API + "/signature/invite", params={"token": tk}).json()
    if nm == "王小明":
        check("簽署頁看得到公司資料", any(p.get("company") == "台灣大哥大股份有限公司" for p in pv["contract"]["parties"]), str(pv)[:150])
    requests.post(API + "/signature/action", json={"token": tk, "action": "sign", "name": nm, "consent": True, "signatureImage": PNG})
pdfr = requests.get(API + f"/contracts/signed-pdf?tenantId={tnc}&contractId={cidc}", headers=HN)
from pypdf import PdfReader as _PR
_txt = "".join(pg.extract_text() or "" for pg in _PR(io.BytesIO(pdfr.content)).pages)
check("已簽署 PDF 含公司名稱、統編、地址、代表人", all(x in _txt for x in ["台灣大哥大股份有限公司", "04595257", "敦化南路", "林之晨", "Acme Corp.", "Jane Doe"]), pdfr.status_code)

# ---- 合約完成後自動寄 Email 通知各方（需本機 SMTP 收信伺服器 127.0.0.1:2525 → /tmp/mails）
import glob, email as _email, time as _t
def _mails(subject_part):
    out = []
    for f in sorted(glob.glob("/tmp/mails/*.eml")):
        m = _email.message_from_bytes(open(f, "rb").read(), policy=__import__("email.policy").policy.default)
        subj = str(m["Subject"])
        if subject_part in subj:
            out.append((m, subj))
    return out
if os.path.isdir("/tmp/mails"):
    stx = requests.get(API + "/contracts/status", headers=HN, params={"tenantId": tnc, "contractId": cidc}).json()["contract"]
    cn = stx.get("completionNotice") or {}
    check("完成時自動寄出通知（2 方＋副本 1 位）", (cn.get("summary") or {}).get("sent") == 3, cn)
    no_c = stx["contractNo"]
    got = _mails(f"【已完成簽署】{no_c}")
    rc_to = sorted(m["To"] for m, _ in got)
    check("收件人為各方與 COMPLETION_CC", rc_to == ["co1@example.com", "co2@example.com", "owner@example.com"], rc_to)
    m0 = got[0][0]
    att = [p for p in m0.walk() if p.get_content_type() == "application/pdf"]
    body = next(p for p in m0.walk() if p.get_content_type() == "text/plain").get_content()
    check("通知信附上已簽署 PDF", len(att) == 1 and att[0].get_payload(decode=True)[:5] == b"%PDF-", [p.get_content_type() for p in m0.walk()])
    check("通知信列出公司資料與下載連結", "台灣大哥大股份有限公司" in body and "/sign/" in body and "SHA-256" in body, body[:300])
    check("同一份合約只自動寄一次", len(got) == 3, len(got))
    rr = requests.post(API + "/contracts/notify-completion", headers=HN, json={"tenantId": tnc, "contractId": cidc}).json()
    check("管理者可手動重寄", rr.get("notice", {}).get("sent") == 3 and len(_mails(f"【已完成簽署】{no_c}")) == 6, rr)
    # 尚未完成的合約不寄、也不能手動寄
    r2 = requests.post(API + "/contracts", headers=HN, json={"title": "未完成合約", "parties": P2}).json()
    inv2 = requests.post(API + "/signature/invite", headers=HN, json={"tenantId": r2["tenantId"], "contractId": r2["contractId"], "expiresAt": exp}).json()
    tk2 = inv2["invites"][0]["previewUrl"].rsplit("/", 1)[1]
    requests.post(API + "/signature/action", json={"token": tk2, "action": "sign", "name": "甲", "consent": True, "signatureImage": PNG})
    no2 = requests.get(API + "/contracts/status", headers=HN, params={"tenantId": r2["tenantId"], "contractId": r2["contractId"]}).json()["contract"]["contractNo"]
    check("部分簽署時不寄完成通知", not _mails(f"【已完成簽署】{no2}"))
    bad3 = requests.post(API + "/contracts/notify-completion", headers=HN, json={"tenantId": r2["tenantId"], "contractId": r2["contractId"]})
    check("未完成的合約不能手動寄完成通知", bad3.status_code in (400, 409), bad3.status_code)
    other = requests.post(API + "/contracts/notify-completion", headers={"Authorization": "Bearer " + tok2}, json={"tenantId": tnc, "contractId": cidc})
    check("別人不能替你的合約寄通知", other.status_code in (403, 404), other.status_code)
ok = sum(1 for _, p in results if p)
print(f"\n結果：{ok}/{len(results)} 通過")
sys.exit(0 if ok == len(results) else 1)
